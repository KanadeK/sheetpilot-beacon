"""Accessible HTML, JSON, and annotated-workbook report exporters."""

from __future__ import annotations

import html
import os
import platform
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sheetpilot_beacon.domain.models import AuditReport, Severity

# Hardened Windows sandboxes can hide platform.machine(), which makes Polars
# report every required flag as "unknown". Bypass only that false-positive
# environment; normal hosts keep Polars' CPU safety check.
_machine_hidden = not platform.machine()
_previous_cpu_check = os.environ.get("POLARS_SKIP_CPU_CHECK")
if _machine_hidden:
    os.environ["POLARS_SKIP_CPU_CHECK"] = "1"
try:
    import polars as pl
finally:
    if _machine_hidden:
        if _previous_cpu_check is None:
            os.environ.pop("POLARS_SKIP_CPU_CHECK", None)
        else:
            os.environ["POLARS_SKIP_CPU_CHECK"] = _previous_cpu_check

SEVERITY_FILL = {
    Severity.CRITICAL: "FECACA",
    Severity.HIGH: "FDE2E2",
    Severity.MEDIUM: "FEF3C7",
    Severity.LOW: "DBEAFE",
    Severity.INFO: "E2E8F0",
}


def _prepare_output(path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    return output


def _excel_safe_text(value: str) -> str:
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _issue_frame(report: AuditReport) -> pl.DataFrame:
    rows = [
        {
            "severity": issue.severity.value,
            "code": issue.code.value,
            "sheet": issue.sheet,
            "cell": issue.cell,
            "title": issue.title,
            "evidence": issue.evidence,
            "recommendation": issue.recommendation,
            "formula": issue.formula_excerpt or "",
        }
        for issue in report.issues
    ]
    return pl.DataFrame(
        rows,
        schema={
            "severity": pl.String,
            "code": pl.String,
            "sheet": pl.String,
            "cell": pl.String,
            "title": pl.String,
            "evidence": pl.String,
            "recommendation": pl.String,
            "formula": pl.String,
        },
    )


def export_json_report(report: AuditReport, output_path: str | Path) -> Path:
    """Write the privacy-minimized report as stable JSON."""

    output = _prepare_output(output_path)
    output.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    return output


def _escape(value: object) -> str:
    return html.escape(str(value), quote=True)


def export_html_report(report: AuditReport, output_path: str | Path) -> Path:
    """Write a standalone responsive and keyboard-readable audit report."""

    output = _prepare_output(output_path)
    frame = _issue_frame(report)
    rows = []
    for issue in frame.iter_rows(named=True):
        location = f"{issue['sheet']}!{issue['cell']}"
        rows.append(
            "<tr>"
            f'<td><span class="badge {_escape(issue["severity"])}">'
            f"{_escape(issue['severity'])}</span></td>"
            f"<td><code>{_escape(issue['code'])}</code></td>"
            f"<td><code>{_escape(location)}</code></td>"
            f"<td><strong>{_escape(issue['title'])}</strong><br>"
            f"<span>{_escape(issue['evidence'])}</span></td>"
            f"<td>{_escape(issue['recommendation'])}</td>"
            f"<td><code>{_escape(issue['formula'])}</code></td>"
            "</tr>"
        )
    severity_cards = "".join(
        f"<li><strong>{_escape(severity.value)}</strong><span>{count}</span></li>"
        for severity, count in report.severity_summary.items()
    )
    external = (
        "<ul>"
        + "".join(f"<li><code>{_escape(link)}</code></li>" for link in report.external_links)
        + "</ul>"
        if report.external_links
        else "<p>None detected.</p>"
    )
    document = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>SheetPilot Beacon · {_escape(report.source_filename)}</title>
  <style>
    :root {{ color-scheme: light; --ink:#102a43; --muted:#52606d;
      --line:#d9e2ec; --paper:#f5f7fa; --teal:#0f766e; }}
    * {{ box-sizing:border-box; }}
    body {{ margin:0; font:15px/1.5 Inter,Segoe UI,Arial,sans-serif;
      color:var(--ink); background:var(--paper); }}
    a.skip {{ position:absolute; left:-9999px; }}
    a.skip:focus {{ left:1rem; top:1rem; background:white; padding:.75rem; z-index:2; }}
    header, main {{ width:min(1440px, calc(100% - 2rem)); margin:auto; }}
    header {{ padding:2.5rem 0 1.25rem; }}
    h1 {{ margin:.2rem 0; font-size:clamp(1.8rem,4vw,3.2rem); }}
    .eyebrow {{ color:var(--teal); font-weight:700; letter-spacing:.08em;
      text-transform:uppercase; }}
    .meta {{ color:var(--muted); overflow-wrap:anywhere; }}
    .cards {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(145px,1fr));
      gap:.75rem; padding:0; list-style:none; }}
    .cards li {{ display:flex; justify-content:space-between; align-items:center;
      background:white; border:1px solid var(--line); border-radius:12px; padding:1rem; }}
    .cards span {{ font-size:1.5rem; font-weight:800; }}
    section {{ background:white; border:1px solid var(--line); border-radius:14px;
      margin:1rem 0; padding:1.1rem; }}
    .table-wrap {{ overflow:auto; max-height:70vh; }}
    table {{ border-collapse:collapse; width:100%; min-width:1050px; }}
    th, td {{ text-align:left; vertical-align:top; padding:.75rem;
      border-bottom:1px solid var(--line); }}
    th {{ position:sticky; top:0; background:#eaf2f8; z-index:1; }}
    code {{ font:13px/1.45 Consolas,monospace; overflow-wrap:anywhere; }}
    .badge {{ display:inline-block; min-width:5rem; text-align:center;
      border:1px solid currentColor; border-radius:999px; padding:.12rem .5rem;
      font-weight:700; }}
    .critical {{ color:#991b1b; background:#fecaca; }}
    .high {{ color:#9f1239; background:#fde2e2; }}
    .medium {{ color:#92400e; background:#fef3c7; }}
    .low {{ color:#1e40af; background:#dbeafe; }}
    .info {{ color:#334155; background:#e2e8f0; }}
    footer {{ color:var(--muted); padding:1rem 0 3rem; }}
  </style>
</head>
<body>
<a class="skip" href="#findings">Skip to findings</a>
<header>
  <div class="eyebrow">SheetPilot Beacon v{_escape(report.tool_version)}</div>
  <h1>Workbook risk audit</h1>
  <p class="meta"><strong>{_escape(report.source_filename)}</strong> ·
    SHA-256 <code>{_escape(report.source_sha256)}</code><br>
    Generated {_escape(report.generated_at.isoformat())} · source unchanged:
    <strong>{"yes" if report.source_unchanged else "no"}</strong></p>
</header>
<main>
  <ul class="cards" aria-label="Issue counts by severity">{severity_cards}</ul>
  <section aria-labelledby="scope-heading">
    <h2 id="scope-heading">Audit scope</h2>
    <p>{report.metrics.sheet_count} sheets,
      {report.metrics.formula_cell_count} formulas,
      {report.metrics.dependency_edge_count} dependency edges,
      {report.metrics.named_range_count} named ranges,
      {report.metrics.merged_range_count} merged ranges.</p>
    <h3>External links</h3>{external}
  </section>
  <section id="findings" aria-labelledby="findings-heading">
    <h2 id="findings-heading">Findings ({len(report.issues)})</h2>
    <div class="table-wrap" tabindex="0" aria-label="Scrollable findings table">
      <table>
        <thead><tr><th>Severity</th><th>Code</th><th>Location</th>
          <th>Evidence</th><th>Recommendation</th><th>Formula excerpt</th></tr></thead>
        <tbody>{"".join(rows)}</tbody>
      </table>
    </div>
  </section>
  <footer>Static analysis only. No workbook values are included in this report.</footer>
</main>
</body>
</html>
"""
    output.write_text(document, encoding="utf-8")
    return output


def _audit_sheet_name(workbook: Any) -> str:
    names = set(workbook.sheetnames)
    base = "Beacon Audit"
    if base not in names:
        return base
    for index in range(2, 100):
        candidate = f"{base} {index}"
        if candidate not in names:
            return candidate
    raise ValueError("workbook already contains too many Beacon Audit sheets")


def export_annotated_workbook(
    source_path: str | Path,
    report: AuditReport,
    output_path: str | Path,
) -> Path:
    """Save findings to a new workbook copy and annotate affected cells."""

    source = Path(source_path)
    output = _prepare_output(output_path)
    if source.resolve() == output.resolve():
        raise ValueError("source and annotated output paths must differ")
    workbook = openpyxl.load_workbook(
        source,
        data_only=False,
        read_only=False,
        keep_links=True,
    )
    try:
        title = _audit_sheet_name(workbook)
        audit = workbook.create_sheet(title, 0)
        audit.sheet_view.showGridLines = False
        audit.freeze_panes = "A7"
        audit.merge_cells("A1:H1")
        audit["A1"] = "SheetPilot Beacon · Workbook Risk Audit"
        audit["A1"].font = Font(size=18, bold=True, color="FFFFFF")
        audit["A1"].fill = PatternFill("solid", fgColor="102A43")
        audit["A1"].alignment = Alignment(vertical="center")
        audit.row_dimensions[1].height = 34
        audit["A3"] = "Source"
        audit["B3"] = _excel_safe_text(report.source_filename)
        audit["A4"] = "SHA-256"
        audit["B4"] = report.source_sha256
        audit["A5"] = "Generated"
        audit["B5"] = report.generated_at.isoformat()
        headers = (
            "Severity",
            "Code",
            "Sheet",
            "Cell",
            "Finding",
            "Evidence",
            "Recommendation",
            "Formula excerpt",
        )
        for column, header in enumerate(headers, start=1):
            cell = audit.cell(row=6, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(wrap_text=True)
        for row, issue in enumerate(report.issues, start=7):
            values = (
                issue.severity.value,
                issue.code.value,
                issue.sheet,
                issue.cell,
                issue.title,
                issue.evidence,
                issue.recommendation,
                issue.formula_excerpt or "",
            )
            for column, value in enumerate(values, start=1):
                audit_cell = audit.cell(
                    row=row,
                    column=column,
                    value=_excel_safe_text(str(value)),
                )
                audit_cell.alignment = Alignment(vertical="top", wrap_text=True)
            audit.cell(
                row=row, column=3
            ).hyperlink = f"#'{issue.sheet.replace(chr(39), chr(39) * 2)}'!{issue.cell}"
            audit.cell(row=row, column=1).fill = PatternFill(
                "solid",
                fgColor=SEVERITY_FILL[issue.severity],
            )
            if issue.sheet in workbook.sheetnames:
                target = workbook[issue.sheet][issue.cell]
                target.fill = PatternFill(
                    "solid",
                    fgColor=SEVERITY_FILL[issue.severity],
                )
                if target.comment is None:
                    target.comment = Comment(
                        f"{issue.title}\n\n{issue.evidence}\n\n{issue.recommendation}",
                        "KanadeK",
                    )
        audit.auto_filter.ref = f"A6:H{max(6, 6 + len(report.issues))}"
        widths = (12, 28, 22, 10, 34, 46, 48, 44)
        for index, width in enumerate(widths, start=1):
            audit.column_dimensions[get_column_letter(index)].width = width
        workbook.save(output)
    finally:
        workbook.close()
    return output
