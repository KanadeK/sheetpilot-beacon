"""Deterministic, licensed workbook fixtures with known audit findings."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from sheetpilot_beacon.domain.models import IssueCode

ExpectedFinding = tuple[IssueCode, str, str]

SAMPLE_EXPECTATIONS: dict[str, tuple[ExpectedFinding, ...]] = {
    "budget_model.xlsx": (
        (IssueCode.HARDCODED_INTRUSION, "Budget", "E6"),
        (IssueCode.BROKEN_REFERENCE, "Budget", "B10"),
        (IssueCode.FORMULA_INCONSISTENCY, "Budget", "D8"),
        (IssueCode.HIGH_FANOUT, "Assumptions", "B2"),
        (IssueCode.HIDDEN_DEPENDENCY, "Budget", "B12"),
    ),
    "sales_forecast.xlsx": (
        (IssueCode.HARDCODED_INTRUSION, "Forecast", "E6"),
        (IssueCode.BROKEN_REFERENCE, "Forecast", "B10"),
        (IssueCode.FORMULA_INCONSISTENCY, "Forecast", "D8"),
        (IssueCode.HIGH_FANOUT, "Inputs", "B2"),
        (IssueCode.EXTERNAL_LINK, "Forecast", "B12"),
    ),
    "engineering_boq.xlsx": (
        (IssueCode.HARDCODED_INTRUSION, "BOQ", "E6"),
        (IssueCode.BROKEN_REFERENCE, "BOQ", "B10"),
        (IssueCode.FORMULA_INCONSISTENCY, "BOQ", "D8"),
        (IssueCode.HIGH_FANOUT, "Assumptions", "B2"),
        (IssueCode.HIDDEN_DEPENDENCY, "BOQ", "B12"),
    ),
}

NAVY = "102A43"
TEAL = "0F766E"
PALE_TEAL = "DDF4F1"
PALE_BLUE = "EAF2F8"
INPUT_BLUE = "0000FF"
FORMULA_BLACK = "000000"
LINK_GREEN = "008000"
EXTERNAL_RED = "FF0000"
WHITE = "FFFFFF"
LIGHT_LINE = Side(style="thin", color="CBD5E1")


@dataclass(frozen=True, slots=True)
class SampleSpec:
    filename: str
    title: str
    input_sheet: str
    model_sheet: str
    named_range: str
    input_label: str
    model_unit: str
    hidden_dependency: bool


SPECS = (
    SampleSpec(
        filename="budget_model.xlsx",
        title="Northwind Operating Budget",
        input_sheet="Assumptions",
        model_sheet="Budget",
        named_range="BudgetGrowth",
        input_label="Annual growth",
        model_unit="USD 000s",
        hidden_dependency=True,
    ),
    SampleSpec(
        filename="sales_forecast.xlsx",
        title="Contoso Sales Forecast",
        input_sheet="Inputs",
        model_sheet="Forecast",
        named_range="SalesGrowth",
        input_label="Pipeline conversion",
        model_unit="Units",
        hidden_dependency=False,
    ),
    SampleSpec(
        filename="engineering_boq.xlsx",
        title="Fabrikam Engineering BOQ",
        input_sheet="Assumptions",
        model_sheet="BOQ",
        named_range="UnitEscalation",
        input_label="Unit-cost escalation",
        model_unit="USD",
        hidden_dependency=True,
    ),
)


def _style_title(sheet: Worksheet, title: str, subtitle: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    title_cell = sheet["A1"]
    title_cell.value = title
    title_cell.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 10
    sheet.merge_cells("A3:F3")
    sheet["A3"] = subtitle
    sheet["A3"].font = Font(color="52606D", italic=True)
    sheet["A3"].alignment = Alignment(wrap_text=True)


def _header(sheet: Worksheet, row: int, values: tuple[str, ...]) -> None:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=LIGHT_LINE)


def _build_cover(workbook: Workbook, spec: SampleSpec) -> None:
    cover = workbook.active
    cover.title = "Cover"
    _style_title(
        cover,
        spec.title,
        "Synthetic workbook for offline formula-risk audit demonstrations",
    )
    rows = (
        ("Version", "1.0"),
        ("Scenario", "Audit exercise"),
        ("Model unit", spec.model_unit),
        ("License", "MIT · synthetic data"),
        ("Privacy", "Contains no real people, companies, or customer records"),
    )
    for row, (label, value) in enumerate(rows, start=5):
        cover.cell(row=row, column=1, value=label).font = Font(bold=True, color=NAVY)
        cover.cell(row=row, column=2, value=value)
    cover.column_dimensions["A"].width = 20
    cover.column_dimensions["B"].width = 58
    for column in ("C", "D", "E", "F"):
        cover.column_dimensions[column].width = 12


def _build_inputs(workbook: Workbook, spec: SampleSpec) -> None:
    sheet = workbook.create_sheet(spec.input_sheet)
    _style_title(sheet, f"{spec.title} · Inputs", "Editable assumptions are blue.")
    _header(sheet, 5, ("Assumption", "Value", "Unit", "Owner"))
    sheet["A6"] = spec.input_label
    sheet["B6"] = 0.08
    sheet["B6"].number_format = "0.0%"
    sheet["B6"].font = Font(color=INPUT_BLUE)
    sheet["B6"].comment = Comment(
        "Synthetic assumption for the SheetPilot Beacon example.",
        "KanadeK",
    )
    sheet["C6"] = "%"
    sheet["D6"] = "Planning"
    sheet["A2"] = "Governed driver"
    sheet["B2"] = 0.05
    sheet["B2"].number_format = "0.0%"
    sheet["B2"].font = Font(color=INPUT_BLUE)
    sheet["B2"].comment = Comment(
        "This governed driver deliberately feeds eight sensitivity cells.",
        "KanadeK",
    )
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 20
    sheet.freeze_panes = "A6"
    workbook.defined_names.add(
        DefinedName(
            spec.named_range,
            attr_text=f"'{spec.input_sheet}'!$B$2",
        )
    )


def _build_model(workbook: Workbook, spec: SampleSpec) -> None:
    sheet = workbook.create_sheet(spec.model_sheet)
    _style_title(
        sheet,
        f"{spec.title} · {spec.model_sheet}",
        "Five deliberate issues are present for a deterministic audit exercise.",
    )
    _header(sheet, 5, ("Metric", "Q1", "Q2", "Q3", "Q4", "Q5"))
    sheet["A6"] = f"Output ({spec.model_unit})"
    for column, input_value in enumerate((100, 110, 120, 130, 140), start=2):
        sheet.cell(row=7, column=column, value=input_value)
        sheet.cell(row=7, column=column).font = Font(color=INPUT_BLUE)
    formulas = {
        "B6": "=B7*1.1",
        "C6": "=C7*1.1",
        "D6": "=D7*1.1",
        "F6": "=F7*1.1",
        "B8": "=B7*2",
        "C8": "=C7*2",
        "D8": "=D7*3",
        "E8": "=E7*2",
        "B10": "=#REF!+B7",
    }
    for coordinate, formula in formulas.items():
        sheet[coordinate] = formula
        sheet[coordinate].font = Font(color=FORMULA_BLACK)
        sheet[coordinate].number_format = "#,##0.0"
    sheet["E6"] = 145.0
    sheet["E6"].font = Font(color=FORMULA_BLACK)
    sheet["E6"].number_format = "#,##0.0"
    sheet["A7"] = "Source inputs"
    sheet["A8"] = "Quality-adjusted output"
    sheet["A10"] = "Legacy subtotal"
    sheet["A12"] = "Linked adjustment"
    if spec.hidden_dependency:
        sheet["B12"] = "='Hidden Inputs'!$B$2*B7"
        sheet["B12"].font = Font(color=LINK_GREEN)
    else:
        sheet["B12"] = "='[legacy_rates.xlsx]Rates'!$B$2*B7"
        sheet["B12"].font = Font(color=EXTERNAL_RED)
    for row in (6, 8, 10, 12):
        for column in range(2, 7):
            sheet.cell(row=row, column=column).number_format = "#,##0.0"
    sheet.freeze_panes = "B6"
    sheet.column_dimensions["A"].width = 30
    for column in range(2, 7):
        sheet.column_dimensions[get_column_letter(column)].width = 14


def _build_sensitivity(workbook: Workbook, spec: SampleSpec) -> None:
    sheet = workbook.create_sheet("Sensitivity")
    _style_title(
        sheet,
        f"{spec.title} · Sensitivity",
        "A governed input intentionally feeds eight downstream formula cells.",
    )
    _header(sheet, 5, ("Case", "Weight", "Driver output"))
    for row in range(6, 14):
        sheet.cell(row=row, column=1, value=f"Case {row - 5}")
        sheet.cell(row=row, column=2, value=row - 5)
        sheet.cell(row=row, column=2).font = Font(color=INPUT_BLUE)
        sheet.cell(row=row, column=3, value=f"={spec.named_range}*1")
        sheet.cell(row=row, column=3).font = Font(color=LINK_GREEN)
        sheet.cell(row=row, column=3).number_format = "0.0%"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 18
    sheet.freeze_panes = "A6"


def _build_checks(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Checks")
    _style_title(
        sheet,
        "Model controls",
        "This training workbook intentionally contains unresolved audit exceptions.",
    )
    _header(sheet, 5, ("Check", "Status", "Where to review", "Notes"))
    checks = (
        ("Formula continuity", "REVIEW", "Model output row", "One seeded constant"),
        ("Reference integrity", "REVIEW", "Legacy subtotal", "One seeded broken link"),
        ("Copy consistency", "REVIEW", "Quality-adjusted row", "One seeded outlier"),
        ("Dependency impact", "REVIEW", "Governed driver", "Eight dependents"),
        ("Link governance", "REVIEW", "Linked adjustment", "Hidden or external source"),
    )
    for row, values in enumerate(checks, start=6):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
        sheet.cell(row=row, column=2).fill = PatternFill("solid", fgColor="FEF3C7")
    for column, width in enumerate((28, 14, 24, 34), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _build_hidden_inputs(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Hidden Inputs")
    sheet["A1"] = "Synthetic hidden adjustment"
    sheet["B2"] = 0.97
    sheet["B2"].number_format = "0.0%"
    sheet.sheet_state = "hidden"


def _build_workbook(spec: SampleSpec, output: Path) -> None:
    workbook = Workbook()
    workbook.properties.creator = "KanadeK"
    workbook.properties.title = spec.title
    workbook.properties.subject = "SheetPilot Beacon deterministic audit sample"
    workbook.properties.description = (
        "Synthetic workbook containing five documented formula-risk findings."
    )
    fixed_time = datetime(2026, 1, 1, 0, 0, 0)
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    _build_cover(workbook, spec)
    _build_inputs(workbook, spec)
    _build_model(workbook, spec)
    _build_sensitivity(workbook, spec)
    _build_checks(workbook)
    _build_hidden_inputs(workbook)
    workbook.active = 0
    workbook.save(output)
    workbook.close()


def generate_sample_workbooks(output_dir: str | Path) -> dict[str, Path]:
    """Generate all three samples with stable data and known findings."""

    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    generated: dict[str, Path] = {}
    for spec in SPECS:
        output = destination / spec.filename
        _build_workbook(spec, output)
        generated[spec.filename] = output
    return generated
