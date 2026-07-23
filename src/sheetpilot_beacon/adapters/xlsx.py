"""Non-evaluating `.xlsx` reader backed by openpyxl."""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

from sheetpilot_beacon.domain.formulas import parse_formula_references
from sheetpilot_beacon.domain.models import (
    CellRecord,
    CellValueKind,
    NamedRangeRecord,
    SheetRecord,
    SheetVisibility,
    WorkbookSnapshot,
)


class WorkbookInputError(ValueError):
    """Raised when an input cannot be safely represented by the MVP."""


@dataclass(frozen=True, slots=True)
class WorkbookReadConfig:
    """Explicit resource limits for workbook ingestion."""

    max_file_bytes: int = 100 * 1024 * 1024
    max_worksheet_cells: int = 1_000_000

    def __post_init__(self) -> None:
        if self.max_file_bytes < 1 or self.max_worksheet_cells < 1:
            raise ValueError("workbook limits must be positive")


def sha256_file(path: Path) -> str:
    """Hash a file without loading it into memory."""

    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _fingerprint(value: object) -> str:
    if isinstance(value, (date, datetime)):
        serializable: object = value.isoformat()
    elif isinstance(value, Decimal):
        serializable = str(value)
    else:
        serializable = value
    payload = json.dumps(
        {"type": type(value).__name__, "value": serializable},
        ensure_ascii=False,
        sort_keys=True,
        default=str,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _cell_kind(value: object, data_type: str) -> CellValueKind:
    if data_type == "e":
        return CellValueKind.ERROR
    if isinstance(value, bool):
        return CellValueKind.BOOLEAN
    if isinstance(value, (datetime, date)):
        return CellValueKind.DATE
    if isinstance(value, (int, float, Decimal)):
        return CellValueKind.NUMBER
    if isinstance(value, str):
        return CellValueKind.TEXT
    return CellValueKind.TEXT


def _defined_names(workbook: Any) -> tuple[NamedRangeRecord, ...]:
    records: list[NamedRangeRecord] = []
    for item in workbook.defined_names.values():
        destinations: list[str] = []
        try:
            for sheet, address in item.destinations:
                quoted = sheet.replace("'", "''")
                destinations.append(f"'{quoted}'!{address}")
        except (AttributeError, TypeError, ValueError):
            destinations = []
        local_sheet = None
        if item.localSheetId is not None:
            index = int(item.localSheetId)
            if 0 <= index < len(workbook.sheetnames):
                local_sheet = str(workbook.sheetnames[index])
        records.append(
            NamedRangeRecord(
                name=str(item.name),
                destinations=tuple(destinations),
                hidden=bool(item.hidden),
                local_sheet=local_sheet,
            )
        )
    return tuple(sorted(records, key=lambda item: item.name.casefold()))


def _visibility(state: str) -> SheetVisibility:
    try:
        return SheetVisibility(state)
    except ValueError as error:
        raise WorkbookInputError(f"unsupported worksheet visibility: {state}") from error


def _validate_input(path: Path, config: WorkbookReadConfig) -> None:
    if not path.exists():
        raise WorkbookInputError(f"workbook does not exist: {path}")
    if not path.is_file():
        raise WorkbookInputError(f"workbook path is not a file: {path}")
    if path.suffix.casefold() != ".xlsx":
        raise WorkbookInputError("v0.1.0 accepts .xlsx files only")
    size = path.stat().st_size
    if size > config.max_file_bytes:
        raise WorkbookInputError(
            f"workbook is {size} bytes; limit is {config.max_file_bytes} bytes"
        )


def read_workbook(
    workbook_path: str | Path,
    *,
    config: WorkbookReadConfig | None = None,
) -> WorkbookSnapshot:
    """Read formulas and structure without calculating or saving the workbook."""

    path = Path(workbook_path)
    effective = config or WorkbookReadConfig()
    _validate_input(path, effective)
    source_hash = sha256_file(path)
    try:
        workbook = openpyxl.load_workbook(
            path,
            data_only=False,
            read_only=False,
            keep_links=True,
        )
    except Exception as error:
        raise WorkbookInputError(f"could not parse workbook: {path.name}") from error

    try:
        sheets: list[SheetRecord] = []
        cells: list[CellRecord] = []
        for worksheet in workbook.worksheets:
            extent = max(1, int(worksheet.max_row)) * max(1, int(worksheet.max_column))
            if extent > effective.max_worksheet_cells:
                raise WorkbookInputError(
                    f"worksheet {worksheet.title!r} spans {extent} cells; "
                    f"limit is {effective.max_worksheet_cells}"
                )
            sheets.append(
                SheetRecord(
                    name=str(worksheet.title),
                    visibility=_visibility(str(worksheet.sheet_state)),
                    max_row=max(1, int(worksheet.max_row)),
                    max_column=max(1, int(worksheet.max_column)),
                    merged_ranges=tuple(
                        sorted(str(cell_range) for cell_range in worksheet.merged_cells.ranges)
                    ),
                )
            )
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                        formula = str(value)
                        if not formula.startswith("="):
                            formula = f"={formula}"
                        cells.append(
                            CellRecord(
                                sheet=str(worksheet.title),
                                coordinate=str(cell.coordinate),
                                value_kind=CellValueKind.FORMULA,
                                formula=formula,
                            )
                        )
                    else:
                        cells.append(
                            CellRecord(
                                sheet=str(worksheet.title),
                                coordinate=str(cell.coordinate),
                                value_kind=_cell_kind(value, str(cell.data_type)),
                                value_fingerprint=_fingerprint(value),
                            )
                        )

        names = _defined_names(workbook)
        name_map = {item.name: item.destinations for item in names}
        external_links: set[str] = set()
        for cell in cells:
            if cell.formula is None:
                continue
            for reference in parse_formula_references(
                cell.formula,
                current_sheet=cell.sheet,
                named_ranges=name_map,
            ):
                if reference.external_book:
                    external_links.add(f"[{reference.external_book}]{reference.sheet}")
        for link in getattr(workbook, "_external_links", ()):
            file_link = getattr(link, "file_link", None)
            target = getattr(file_link, "Target", None)
            if target:
                external_links.add(str(target))

        return WorkbookSnapshot(
            filename=path.name,
            sha256=source_hash,
            sheets=tuple(sheets),
            cells=tuple(cells),
            named_ranges=names,
            external_links=tuple(sorted(external_links)),
        )
    finally:
        workbook.close()
