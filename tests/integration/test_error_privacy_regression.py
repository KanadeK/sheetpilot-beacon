from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from sheetpilot_beacon.adapters.exporters import (
    export_annotated_workbook,
    export_html_report,
    export_json_report,
)
from sheetpilot_beacon.adapters.samples import generate_sample_workbooks
from sheetpilot_beacon.cli import main
from sheetpilot_beacon.services.audit import audit_workbook
from sheetpilot_beacon.ui import audit_upload


@pytest.mark.integration
def test_cli_turns_output_permission_error_into_actionable_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    source = generate_sample_workbooks(tmp_path / "samples")["budget_model.xlsx"]
    denied_output = tmp_path / "denied"
    original_mkdir = Path.mkdir

    def deny_output_directory(path: Path, *args: object, **kwargs: object) -> None:
        if path == denied_output:
            raise PermissionError("permission denied by regression fixture")
        original_mkdir(path, *args, **kwargs)

    monkeypatch.setattr(Path, "mkdir", deny_output_directory)

    exit_code = main(["scan", str(source), "--output-dir", str(denied_output)])

    assert exit_code == 1
    assert "permission denied by regression fixture" in capsys.readouterr().err


@pytest.mark.integration
def test_reports_redact_credentials_and_omit_raw_constant_values(tmp_path: Path) -> None:
    source = tmp_path / "privacy.xlsx"
    secret_token = "sk-ABCDEFGHIJKLMNOPQRSTUV"
    private_value = "private-customer-value-8742"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Model"
    worksheet["A1"] = f'="api_key={secret_token}"&#REF!'
    worksheet["B1"] = private_value
    workbook.save(source)
    workbook.close()

    report = audit_workbook(source)
    json_path = export_json_report(report, tmp_path / "privacy.audit.json")
    html_path = export_html_report(report, tmp_path / "privacy.audit.html")
    annotated_path = export_annotated_workbook(
        source,
        report,
        tmp_path / "privacy.annotated.xlsx",
    )

    json_text = json_path.read_text(encoding="utf-8")
    html_text = html_path.read_text(encoding="utf-8")
    assert secret_token not in json_text
    assert secret_token not in html_text
    assert private_value not in json_text
    assert private_value not in html_text
    assert "[REDACTED]" in json_text

    annotated = openpyxl.load_workbook(annotated_path, data_only=False)
    try:
        audit_sheet = annotated["Beacon Audit"]
        excerpts = [
            str(audit_sheet.cell(row=row, column=8).value)
            for row in range(7, audit_sheet.max_row + 1)
        ]
        assert any("[REDACTED]" in excerpt for excerpt in excerpts)
        assert all(secret_token not in excerpt for excerpt in excerpts)
        assert annotated["Model"]["B1"].value == private_value
    finally:
        annotated.close()


@pytest.mark.integration
def test_annotated_export_escapes_formula_like_sheet_names_and_avoids_title_collision(
    tmp_path: Path,
) -> None:
    source = tmp_path / "formula-injection.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "=CMD"
    worksheet["A1"] = "=#REF!"
    workbook.create_sheet("Beacon Audit")
    workbook.save(source)
    workbook.close()

    report = audit_workbook(source)
    output = export_annotated_workbook(source, report, tmp_path / "annotated.xlsx")

    annotated = openpyxl.load_workbook(output, data_only=False)
    try:
        assert annotated.sheetnames[0] == "Beacon Audit 2"
        location_sheet = annotated["Beacon Audit 2"].cell(row=7, column=3)
        assert location_sheet.data_type == "s"
        assert location_sheet.value == "'=CMD"
    finally:
        annotated.close()


@pytest.mark.parametrize(
    ("filename", "payload", "message"),
    [
        ("empty.xlsx", b"", "empty"),
        ("unsafe.csv", b"data", ".xlsx filename"),
    ],
)
def test_upload_boundary_rejects_invalid_payloads(
    tmp_path: Path,
    filename: str,
    payload: bytes,
    message: str,
) -> None:
    with pytest.raises(ValueError, match=message):
        audit_upload(filename, payload, work_dir=tmp_path)


def test_privacy_report_json_is_valid_and_contains_only_fingerprints(tmp_path: Path) -> None:
    source = tmp_path / "values.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "sensitive-row-value"
    worksheet["B1"] = "=A1"
    workbook.save(source)
    workbook.close()

    payload = json.loads(audit_workbook(source).model_dump_json())

    assert "sensitive-row-value" not in json.dumps(payload)
    assert payload["metrics"]["populated_cell_count"] == 2
