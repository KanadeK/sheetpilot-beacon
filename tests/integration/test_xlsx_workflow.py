from __future__ import annotations

import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pytest

from sheetpilot_beacon.adapters.exporters import (
    export_annotated_workbook,
    export_html_report,
    export_json_report,
)
from sheetpilot_beacon.adapters.samples import SAMPLE_EXPECTATIONS, generate_sample_workbooks
from sheetpilot_beacon.adapters.xlsx import read_workbook
from sheetpilot_beacon.domain.models import IssueCode, SheetVisibility
from sheetpilot_beacon.services.audit import FixedClock, audit_workbook


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(64 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@pytest.mark.integration
def test_sample_generator_creates_structurally_rich_workbooks(tmp_path: Path) -> None:
    paths = generate_sample_workbooks(tmp_path)

    assert set(paths) == {
        "budget_model.xlsx",
        "engineering_boq.xlsx",
        "sales_forecast.xlsx",
    }
    budget = read_workbook(paths["budget_model.xlsx"])
    sales = read_workbook(paths["sales_forecast.xlsx"])
    assert any(sheet.visibility is SheetVisibility.HIDDEN for sheet in budget.sheets)
    assert any(sheet.merged_ranges for sheet in budget.sheets)
    assert {item.name for item in budget.named_ranges} >= {"BudgetGrowth"}
    assert sales.external_links == ("[legacy_rates.xlsx]Rates",)


@pytest.mark.integration
def test_all_seeded_findings_are_detected_without_extra_findings(tmp_path: Path) -> None:
    paths = generate_sample_workbooks(tmp_path)
    clock = FixedClock(datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC))

    for filename, expected in SAMPLE_EXPECTATIONS.items():
        report = audit_workbook(paths[filename], clock=clock)
        observed = {(issue.code, issue.sheet, issue.cell) for issue in report.issues}

        assert observed == set(expected)
        assert report.source_unchanged is True
        assert report.generated_at == clock.now()
        assert sum(report.issue_summary.values()) == 5


@pytest.mark.integration
def test_audit_exports_real_reports_and_never_changes_source(tmp_path: Path) -> None:
    source = generate_sample_workbooks(tmp_path / "samples")["budget_model.xlsx"]
    before_hash = file_hash(source)
    report = audit_workbook(
        source,
        clock=FixedClock(datetime(2026, 1, 2, 3, 4, 5, tzinfo=UTC)),
    )
    output_dir = tmp_path / "reports"
    html_path = export_html_report(report, output_dir / "budget-audit.html")
    json_path = export_json_report(report, output_dir / "budget-audit.json")
    xlsx_path = export_annotated_workbook(
        source,
        report,
        output_dir / "budget-annotated.xlsx",
    )

    assert file_hash(source) == before_hash
    assert html_path.read_text(encoding="utf-8").count("<tr") >= 6
    json_payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert json_payload["source_sha256"] == before_hash
    assert len(json_payload["issues"]) == 5

    source_book = openpyxl.load_workbook(source, data_only=False)
    annotated_book = openpyxl.load_workbook(xlsx_path, data_only=False)
    try:
        assert "Beacon Audit" not in source_book.sheetnames
        assert annotated_book.sheetnames[0] == "Beacon Audit"
        assert annotated_book["Budget"]["B6"].value == source_book["Budget"]["B6"].value
        assert annotated_book["Budget"]["E6"].comment is not None
    finally:
        source_book.close()
        annotated_book.close()


@pytest.mark.integration
def test_annotated_export_rejects_overwriting_source(tmp_path: Path) -> None:
    source = generate_sample_workbooks(tmp_path)["engineering_boq.xlsx"]
    report = audit_workbook(source)

    with pytest.raises(ValueError, match="must differ"):
        export_annotated_workbook(source, report, source)


def test_expected_manifest_contract_has_exactly_fifteen_findings() -> None:
    all_expected = [
        finding
        for workbook_findings in SAMPLE_EXPECTATIONS.values()
        for finding in workbook_findings
    ]

    assert len(all_expected) == 15
    assert {finding[0] for finding in all_expected} >= {
        IssueCode.HARDCODED_INTRUSION,
        IssueCode.BROKEN_REFERENCE,
        IssueCode.FORMULA_INCONSISTENCY,
        IssueCode.HIGH_FANOUT,
        IssueCode.HIDDEN_DEPENDENCY,
        IssueCode.EXTERNAL_LINK,
    }
