from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
EXAMPLES = REPOSITORY_ROOT / "examples"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _run_cli(*arguments: object) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        (
            sys.executable,
            "-m",
            "sheetpilot_beacon.cli",
            *(str(argument) for argument in arguments),
        ),
        cwd=REPOSITORY_ROOT,
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )


@pytest.mark.e2e
def test_scan_command_audits_real_workbook_and_exports_three_formats(tmp_path: Path) -> None:
    source = EXAMPLES / "budget_model.xlsx"
    source_hash = _sha256(source)
    output = tmp_path / "scan"

    completed = _run_cli("scan", source, "--output-dir", output)

    assert completed.returncode == 0, completed.stderr
    assert "5 findings, 18 formulas, 19 dependency edges" in completed.stdout
    assert _sha256(source) == source_hash
    payload = json.loads((output / "budget_model.audit.json").read_text(encoding="utf-8"))
    assert len(payload["issues"]) == 5
    assert payload["source_unchanged"] is True
    assert (
        (output / "budget_model.audit.html")
        .read_text(encoding="utf-8")
        .startswith("<!doctype html>")
    )
    workbook = openpyxl.load_workbook(output / "budget_model.annotated.xlsx")
    try:
        assert workbook.sheetnames[0] == "Beacon Audit"
        assert workbook["Budget"]["E6"].comment is not None
    finally:
        workbook.close()


@pytest.mark.e2e
def test_diff_command_pinpoints_formula_replaced_by_constant(tmp_path: Path) -> None:
    before = EXAMPLES / "budget_model.xlsx"
    after = tmp_path / "budget_model_revised.xlsx"
    workbook = openpyxl.load_workbook(before, data_only=False)
    try:
        workbook["Budget"]["C6"] = 121.0
        workbook.save(after)
    finally:
        workbook.close()

    output = tmp_path / "diff"
    completed = _run_cli("diff", before, after, "--output-dir", output)

    assert completed.returncode == 0, completed.stderr
    assert "1 changes" in completed.stdout
    payload = json.loads(
        (output / "budget_model-to-budget_model_revised.drift.json").read_text(encoding="utf-8")
    )
    assert [(change["kind"], change["sheet"], change["cell"]) for change in payload["changes"]] == [
        ("formula_replaced_by_constant", "Budget", "C6")
    ]
    assert "formula_replaced_by_constant" in (
        output / "budget_model-to-budget_model_revised.drift.html"
    ).read_text(encoding="utf-8")


@pytest.mark.e2e
def test_demo_command_audits_all_samples_and_builds_static_index(tmp_path: Path) -> None:
    output = tmp_path / "demo"

    completed = _run_cli(
        "demo",
        "--samples-dir",
        EXAMPLES,
        "--output-dir",
        output,
    )

    assert completed.returncode == 0, completed.stderr
    assert "audited 3 workbooks and found 15 seeded issues" in completed.stdout
    reports = sorted(output.glob("*.audit.json"))
    assert len(reports) == 3
    assert (
        sum(len(json.loads(report.read_text(encoding="utf-8"))["issues"]) for report in reports)
        == 15
    )
    index = (output / "index.html").read_text(encoding="utf-8")
    assert index.count("Open audit report") == 3
    assert "fifteen seeded findings" in index


@pytest.mark.e2e
@pytest.mark.parametrize("failure_kind", ["missing", "wrong-extension", "corrupt"])
def test_scan_command_reports_input_failures(
    tmp_path: Path,
    failure_kind: str,
) -> None:
    if failure_kind == "missing":
        source = tmp_path / "missing.xlsx"
        expected = "No such file or directory"
    elif failure_kind == "wrong-extension":
        source = tmp_path / "workbook.csv"
        source.write_text("not,xlsx", encoding="utf-8")
        expected = "accepts .xlsx files only"
    else:
        source = tmp_path / "corrupt.xlsx"
        source.write_bytes(b"not an OOXML archive")
        expected = "could not parse workbook"

    completed = _run_cli("scan", source, "--output-dir", tmp_path / "output")

    assert completed.returncode == 1
    assert "sheetpilot: error:" in completed.stderr
    assert expected.casefold() in completed.stderr.casefold()
